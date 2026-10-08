"""Build the Airlite Product Price Template from the ECO template structure.
Usage: python -I build_airlite.py ECO_TEMPLATE.xlsx AIRLITE_CONFIGURATOR.xlsm OUT.xlsx
"""
import sys, copy, warnings
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.comments import Comment

warnings.filterwarnings("ignore")
eco_path, cfg_path, out_path = sys.argv[1:4]

PFX, PCODE = "AIR", "AIR"
PNAME = "Airlite Smoke and Heat Ventilator AIR"
CLASSIF = f"{PNAME} - Product Classification"
YELLOW = PatternFill("solid", fgColor="FFFF00")

# ---------------------------------------------------------------- source data
cfg = openpyxl.load_workbook(cfg_path, data_only=True)["Input"]
desc_en, desc_de = {}, {}
for r in range(1, 130):
    k = cfg.cell(r, 58).value                      # BF code, BG EN, BH code, BI DE
    if k not in (None, "") and cfg.cell(r, 59).value:
        desc_en.setdefault(str(k), []).append((r, cfg.cell(r, 59).value))
        desc_de.setdefault(str(k), []).append((r, cfg.cell(r, 61).value))


def lbl(code, rows):
    """EN/DE label for a code within a BF row band (codes like 'S' repeat)."""
    lo, hi = rows
    en = next(v for r, v in desc_en[code] if lo <= r <= hi)
    de = next(v for r, v in desc_de[code] if lo <= r <= hi)
    return en, de


def dropdown(row, cells):
    """Selectable values exactly as the configurator's data-validation list."""
    return [cfg.cell(row, c).value for c in cells if cfg.cell(row, c).value not in (None, "", 0)]


M = 13  # column M
versions = dropdown(13, range(M, M + 2))
lengths = dropdown(15, range(M, M + 9))
louvres = dropdown(19, range(M, M + 6))
controls = dropdown(20, range(M, M + 13))
releases = dropdown(21, range(M, M + 6))
brackets = dropdown(22, range(M, M + 2))
bases = dropdown(23, range(M, M + 3))
flanges = dropdown(24, range(M, M + 20))
coatings = dropdown(25, range(M, M + 6))
thick = dropdown(27, range(M, M + 3))
installs = dropdown(28, range(M, M + 4))
baffles = dropdown(30, range(M, M + 2))
deliveries = dropdown(31, range(M, M + 2))

# picklist code -> list of (value, label EN, label DE)
PL = {}
PL["Version"] = [(v, *lbl(v, (5, 6))) for v in versions]
PL["Length"] = [(str(v), f"{v} mm", f"{v} mm") for v in lengths]
PL["Insulation"] = [("I", *lbl("I", (9, 9)))]
PL["Louvre"] = [(v, *lbl(v, (12, 17))) for v in louvres]
PL["Controls"] = [(v, *lbl(v, (20, 33))) for v in controls]
PL["Thermal Release"] = [(v, *lbl(v, (74, 79))) for v in releases]
PL["Control Bracket"] = [(v, *lbl(v, (39, 40))) for v in brackets]
PL["Base Design"] = [(v, *lbl(v, (42, 44))) for v in bases]
PL["Flange Type"] = [(v, *lbl(v, (46, 65))) for v in flanges]
PL["Coating"] = [(v, *lbl(v, (67, 72))) for v in coatings]
PL["Coating Thickness"] = [(str(v), f"{v} µm", f"{v} µm") for v in thick]
PL["Installation Type"] = [(v, *lbl(v, (88, 91))) for v in installs]
PL["Wind Baffle"] = [(v, *lbl(v, (84, 85))) for v in baffles]
PL["Delivery"] = [(v, *lbl(v, (94, 95))) for v in deliveries]
PL["Ambient Temp"] = [("T(-15)", "Ambient air temperature class T(-15)", "Umgebungstemperatur T(-15)"),
                      ("T(-25)", "Ambient air temperature class T(-25)", "Umgebungstemperatur T(-25)")]

PL_DESC = {  # picklist code suffix -> (description, German name)
    "Version": ("Device version (EN / Standard)", "Geräteausführung"),
    "Length": ("Opening length in mm (300 mm louvre steps)", "Öffnungslänge"),
    "Insulation": ("Insulation", "Isolation"),
    "Louvre": ("Louvre / filling type", "Lamellentyp"),
    "Controls": ("Actuation / drive type", "Antrieb"),
    "Thermal Release": ("Thermal release / fire override", "Thermische Auslösung"),
    "Control Bracket": ("Control bracket design", "Antriebskonsole"),
    "Base Design": ("Base design", "Basisausbildung"),
    "Flange Type": ("Flange type", "Flanschausbildung"),
    "Coating": ("Coating / paint system", "Beschichtung"),
    "Coating Thickness": ("Coating thickness in µm", "Schichtdicke"),
    "Installation Type": ("Installation situation and side-wind influence", "Einbausituation"),
    "Wind Baffle": ("Wind baffle", "Windleitwand"),
    "Delivery": ("Delivery state", "Lieferzustand"),
    "Ambient Temp": ("Ambient air temperature class", "Umgebungstemperatur"),
}

# ------------------------------------------------------------- attributes
# (Name ENG, Label, Description, DataType, IsRequired, Default, picklist suffix,
#  code suffix, Name GER, category, P&A default)
CFG_ATTRS = [
    ("Version", "Version", "Airlite version: 1E = EN 12101-2 certified, 1S = standard", "Picklist", "TRUE", "1E", "Version", "Version", "Geräteausführung", "Device Configuration"),
    ("Width", "Width (mm)", "Opening width in mm. Min. 500, max. 2000; min. 700 with 2 actuators", "Number", "TRUE", None, None, "Width", "Öffnungsbreite", "Device Configuration"),
    ("Length", "Length (mm)", "Opening length in mm, 300 mm louvre steps", "Picklist", "TRUE", None, "Length", "Length", "Öffnungslänge", "Device Configuration"),
    ("Flange Width", "Flange Width (mm)", "Outer flange width in mm. Default = width + 250 mm (N5); limits depend on flange type and base design", "Number", "FALSE", None, None, "FlangeWidth", "Äußere Flanschbreite", "Base & Flange"),
    ("Flange Length", "Flange Length (mm)", "Outer flange length in mm. Default = length + 250 mm (N5); limits depend on flange type and base design", "Number", "FALSE", None, None, "FlangeLength", "Äußere Flanschlänge", "Base & Flange"),
    ("Insulation", "Insulation", "Insulation (only option in configurator: I)", "Picklist", "TRUE", "I", "Insulation", "Insulation", "Isolation", "Device Configuration"),
    ("Louvre", "Louvre", "Louvre / filling type", "Picklist", "TRUE", None, "Louvre", "Louvre", "Lamelle", "Filling"),
    ("Controls", "Controls", "Pneumatic cylinder (P) or electric actuator (M), quantity and size", "Picklist", "TRUE", None, "Controls", "Controls", "Antrieb", "Control & Drive"),
    ("Thermal Release", "Thermal Release", "Thermal release / failsafe option", "Picklist", "TRUE", "FX", "Thermal Release", "ThermalRelease", "Thermische Auslösung", "Control & Drive"),
    ("Control Bracket Design", "Control Bracket Design", "Long or short control bracket", "Picklist", "TRUE", None, "Control Bracket", "ControlBracket", "Antriebskonsole", "Control & Drive"),
    ("Base Design", "Base Design", "Base construction", "Picklist", "TRUE", None, "Base Design", "BaseDesign", "Basisausbildung", "Base & Flange"),
    ("Flange Type", "Flange Type", "Flange / fixing detail", "Picklist", "TRUE", "N5", "Flange Type", "FlangeType", "Flanschausbildung", "Base & Flange"),
    ("Coating", "Coating", "Coating / paint system", "Picklist", "TRUE", "X", "Coating", "Coating", "Beschichtung", "Finish"),
    ("RAL Colour", "RAL Colour", "RAL colour number (required when coating is not X)", "Text", "FALSE", None, None, "RALColour", "RAL-Farbe", "Finish"),
    ("Coating Thickness", "Coating Thickness (µm)", "Coating thickness in µm", "Picklist", "FALSE", "60", "Coating Thickness", "CoatingThickness", "Schichtdicke", "Finish"),
    ("Installation Type", "Installation Type", "Flat roof (A0/A2) or wall (E1/E2), with/without upstand", "Picklist", "TRUE", None, "Installation Type", "InstallationType", "Einbauart", "Installation"),
    ("Mounting Angle", "Mounting Angle (°)", "Mounting angle to horizontal, 0–90°", "Number", "TRUE", None, None, "MountingAngle", "Einbauwinkel", "Installation"),
    ("Wind Baffle", "Wind Baffle", "Movable wind baffle (WM) or none (W0)", "Picklist", "TRUE", None, "Wind Baffle", "WindBaffle", "Windleitwand", "Installation"),
    ("Delivery", "Delivery", "Completely assembled or louvres delivered separately", "Picklist", "TRUE", "CA", "Delivery", "Delivery", "Lieferzustand", "Options"),
    ("Snow Load", "Snow Load SL (N/m²)", "Design snow load in N/m²", "Number", "TRUE", None, None, "SnowLoad", "Schneelast", "Loads & Pressure"),
    ("Wind Load", "Wind Load (N/m²)", "Design wind load in N/m²", "Number", "TRUE", None, None, "WindLoad", "Windlast", "Loads & Pressure"),
    ("Ambient Air Temperature", "Ambient Air Temperature", "T(-25) with 24 V actuators (M1B24/M2B24), otherwise T(-15)", "Picklist", "TRUE", "T(-15)", "Ambient Temp", "AmbientTemp", "Umgebungstemperatur", "Loads & Pressure"),
    ("Pressure Day to Day", "Pressure Day to Day (bar)", "Pneumatic pressure for daily ventilation, min. 6 bar (pneumatic controls only)", "Number", "FALSE", "6", None, "PressureDaily", "Druck Lüftung", "Loads & Pressure"),
    ("Pressure Smoke Ventilation", "Pressure Smoke Ventilation (bar)", "Pneumatic pressure for smoke ventilation, min. 6 bar (pneumatic controls only)", "Number", "FALSE", "6", None, "PressureSmoke", "Druck RWA", "Loads & Pressure"),
]
OUT_ATTRS = [  # calculated outputs (read-only); source cell in configurator
    ("Cylinder", "Cylinder", "Pneumatic cylinder type. Source: SL Pneumatic Cylinders!B82", "Text", "CylinderType", "Zylinder"),
    ("CO2 Bottle", "CO2 Bottle (g)", "CO2 bottle size in g. Source: SL Pneumatic Cylinders!G69", "Number", "CO2Bottle", "CO2-Flasche"),
    ("Pressure Standby Reservoir", "Pressure Standby Reservoir (ltr.)", "Pressure standby reservoir. Source: SL Pneumatic Cylinders!K51", "Text", "StandbyReservoir", "Druckluftbehälter"),
    ("Electric Drive", "Electric Drive", "Electric drive check. Source: SL Electric Drive!B88", "Text", "ElectricDrive", "Elektroantrieb"),
    ("Wind Suction Load WL", "Wind Suction Load WL (N/m²)", "Wind suction load. Source: Windload!B31", "Number", "WindSuctionLoad", "Windsoglast"),
    ("Geometric Area Av", "Geometric Area Av (m²)", "Width × Length / 10^6. Source: Input!C47", "Number", "GeomAreaAv", "Geometrische Fläche Av"),
    ("Aerodynamic Area Aa", "Aerodynamic Area Aa (m²)", "Aerodynamic area incl. side wind. Source: Aerodynamik!B78", "Number", "AeroAreaAa", "Aerodynamische Fläche Aa"),
    ("U-Value", "U-Value (W/(m²K))", "U-value of unit. Source: U-values!O33", "Number", "UValue", "U-Wert"),
    ("Louvre Weight", "Weight of One Louvre (kg)", "Source: Louvre Variants!C30", "Number", "LouvreWeight", "Gewicht einer Lamelle"),
    ("Total Weight", "Total Weight (kg)", "Total unit weight. Source: Weights!J96", "Number", "TotalWeight", "Gesamtgewicht"),
    ("Max SL Pneumatic Cylinder", "Max. SL Pneumatic Cylinder (N/m²)", "Source: SL Pneumatic Cylinders!B71", "Number", "MaxSLCylinder", "Max. SL Zylinder"),
    ("Max SL Motor", "Max. SL Motor (N/m²)", "Source: SL Electric Drive", "Number", "MaxSLMotor", "Max. SL Motor"),
]

# ------------------------------------------------------------- constraints
C = lambda s: f"{PFX}_{s}"
P2 = [c for c in controls if c[1] == "2"]
PNEU = [c for c in controls if c.startswith("P")]
ELEC = [c for c in controls if c.startswith("M")]
J = "  ".join
CONSTRAINTS = [
    ("ALL", C("Width"), "500 – 2000", "Opening width min. 500 mm, max. 2000 mm (both versions)"),
    *[(c, C("Width"), "700 – 2000", "Min. width with 2 actuators is 700 mm") for c in P2],
    *[(c, C("ThermalRelease"), "FX", "Electric controls only in combination with FX") for c in ELEC],
    *[(c, C("ThermalRelease"), "FX  FS68  FS93  FF68  FF93", "1E only: pneumatic controls not with FFX (1S allows FFX)") for c in PNEU],
    ("GL24", C("InstallationType"), "A0  A2", "GL24 is invalid by installation into wall (E1, E2)"),
    ("GL24", C("MountingAngle"), "0 – 20", "1E only: GL24 max. mounting angle 20°"),
    ("E", C("FlangeType"), "N5", "Base design E is only possible with a N5 flange"),
    ("PG  PS  PM  W  S", C("RALColour"), "required", "Insert RAL colour when coating is not X"),
    ("PG  PS  PM  S", f"{C('Width')} / {C('Length')}", "shorter side ≤ 1900", "Attention: the shorter side must not be larger than 1900 mm for coated units"),
    ("A0", C("WindBaffle"), "WM", "Flat roof installation requires movable wind baffle"),
    ("A2", C("WindBaffle"), "WM", "Flat roof installation requires movable wind baffle"),
    ("E1", C("WindBaffle"), "W0", "No movable wind baffle required for wall installation"),
    ("E2", C("WindBaffle"), "W0", "No movable wind baffle required for wall installation"),
    ("E1", C("MountingAngle"), "90", "1E only: installation in wall only with mounting angle 90°"),
    ("E2", C("MountingAngle"), "90", "1E only: installation in wall only with mounting angle 90°"),
    ("ALL", C("MountingAngle"), "0 – 90", "Mounting angle 0–90°"),
    *[(l, C("Delivery"), "CA", "1E only: break down (BD) only in combination with glass louvres (GL24)") for l in louvres if l != "GL24"],
    *[(c, C("AmbientTemp"), "T(-25)", "24 V actuators: ambient air temperature T(-25)") for c in ("M1B24", "M2B24")],
    *[(c, C("AmbientTemp"), "T(-15)", "All other controls: ambient air temperature T(-15)") for c in controls if c not in ("M1B24", "M2B24")],
    ("P* (pneumatic)", f"{C('PressureDaily')} / {C('PressureSmoke')}", "6 – max. cylinder pressure (≤ 30)", "Pressure out of the valid pressure range"),
    ("FFX  FF68  FF93", f"{C('PressureDaily')} / {C('PressureSmoke')}", "equal, ≤ 12", "Failsafe: day-to-day pressure = smoke pressure, max. 12 bar"),
    ("N1-16  N1-24", f"{C('FlangeWidth')} / {C('FlangeLength')}", "W +150…+200 / L +150…+500", "Flange dimension limits (Flanges sheet)"),
    ("N2  N3  N4  N7  NS", f"{C('FlangeWidth')} / {C('FlangeLength')}", "W ≥ +150 / L ≥ +150", "Flange dimension limits (Flanges sheet)"),
    ("N5", f"{C('FlangeWidth')} / {C('FlangeLength')}", "+150…+500 (base E: exactly +250)", "Flange dimension limits (Flanges sheet)"),
    ("N6 - *", f"{C('FlangeWidth')} / {C('FlangeLength')}", "S: W ≥ +300 / L ≥ +355; H: W ≥ +250 / L ≥ +330; E: ≥ +300", "Flange dimension limits (Flanges sheet)"),
    ("N8 - *", f"{C('FlangeWidth')} / {C('FlangeLength')}", "S: W ≥ +350 / L ≥ +380; H: W ≥ +250 / L ≥ +330; E: ≥ +300", "Flange dimension limits (Flanges sheet)"),
    ("Calculated", C("SnowLoad"), "≤ max. SL of cylinder / motor", "Reduce snow load, change dimensions or change cylinder type (SL Pneumatic Cylinders / SL Electric Drive tables)"),
    ("Calculated (1E)", C("WindLoad"), "WL check OK", "Wind suction load too high (Windload sheet)"),
    ("Calculated", C("ThermalRelease"), "CO2 bottle ≥ 20 g", "Pneumatic with thermal release needs CO2 bottle ≥ 20 g"),
]
DRIVER = {}  # infer driver attribute from the driver value
for a, vals in [("Controls", controls), ("Louvre", louvres), ("BaseDesign", ["E"]),
                ("Coating", coatings), ("InstallationType", installs), ("ThermalRelease", releases),
                ("FlangeType", flanges)]:
    for v in vals:
        DRIVER.setdefault(v, C(a))
DRIVER.update({"PG  PS  PM  W  S": C("Coating"), "PG  PS  PM  S": C("Coating"),
               "P* (pneumatic)": C("Controls"), "FFX  FF68  FF93": C("ThermalRelease"),
               "N1-16  N1-24": C("FlangeType"), "N2  N3  N4  N7  NS": C("FlangeType"),
               "N6 - *": C("FlangeType"), "N8 - *": C("FlangeType"),
               "ALL": "—", "Calculated": "—", "Calculated (1E)": C("Version")})

# ------------------------------------------------------------- workbook
wb = openpyxl.load_workbook(eco_path)


def clear(ws, first):
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= first:
            ws.unmerge_cells(str(mr))
    for row in ws.iter_rows(min_row=first, max_row=max(ws.max_row, first)):
        for c in row:
            c.value = None
            c.comment = None


def put(ws, first, rows, style_row=None, ncols=None):
    """Write rows starting at `first`, copying cell styles from style_row."""
    style_row = style_row or first
    ncols = ncols or max(len(r) for r in rows)
    styles = [copy.copy(ws.cell(style_row, c)._style) for c in range(1, ncols + 1)]
    for i, r in enumerate(rows):
        for j in range(ncols):
            cell = ws.cell(first + i, j + 1)
            cell._style = copy.copy(styles[j])
            cell.value = r[j] if j < len(r) else None


# Product
ws = wb["Product"]
prod = [None] * 23
prod[3:23] = ["yes", PNAME, None, "Allowed", PCODE, None, None, "TRUE",
              "Natural smoke and heat exhaust ventilator (NSHEV) with louvres. Pneumatic or electric control. Certified according to EN 12101-2 (version 1E).",
              "Airlite – configurable via RCA. Version 1E/1S. Opening width 500–2000 mm, length 1200–3600 mm in 300 mm louvre steps.",
              None, None, None, None, None, "TRUE", "FALSE", None, "Natural Ventilation", "Global"]
style_src = 6
clear(ws, 6)
put(ws, 6, [prod], style_src, 23)
ws["N6"].fill = YELLOW; ws["N6"].comment = Comment("Fill in: product owner e-mail", "Claude")
ws["P6"].fill = YELLOW; ws["P6"].comment = Comment("Fill in: availability date (DD.MM.YYYY)", "Claude")

# AttributeDefinition
ws = wb["AttributeDefinition"]
rows, seq = [], 10
for name, label, desc, dt, req, dflt, pl, code, ger, cat in CFG_ATTRS:
    rows.append(["X", name, label, desc, str(seq), dt, "TRUE", req, dflt,
                 f"{PFX} {pl}" if pl else None, C(code), C(code), ger])
    seq += 10
out_first = len(rows)
seq = 300
for name, label, desc, dt, code, ger in OUT_ATTRS:
    rows.append(["X", name, label, f"Calculated (read-only). {desc}", str(seq), dt, "TRUE", "FALSE",
                 None, None, C(code), C(code), ger])
    seq += 10
clear(ws, 3)
put(ws, 3, rows, 3, 13)

# Product and Attributes
ws = wb["Product and Attributes"]
rows, seq = [], 10
for name, label, desc, dt, req, dflt, pl, code, ger, cat in CFG_ATTRS:
    vals = J(v for v, _, _ in PL[pl]) if pl else None
    rows.append([PCODE, str(seq), C(code), dflt, cat, vals, None, CLASSIF])
    seq += 10
seq = 300
for name, label, desc, dt, code, ger in OUT_ATTRS:
    rows.append([PCODE, str(seq), C(code), None, "Calculated Outputs", None, None, CLASSIF])
    seq += 10
clear(ws, 3)
put(ws, 3, rows, 3, 8)

# AttributePicklist
ws = wb["AttributePicklist"]
used = [a[6] for a in CFG_ATTRS if a[6]]
rows = [[f"{PFX} {pl}", f"{PFX} {pl}", PL_DESC[pl][0], "Active", "Text", None, PL_DESC[pl][1]] for pl in used]
clear(ws, 3)
put(ws, 3, rows, 3, 7)

# AttributePicklistValue
ws = wb["AttributePicklistValue"]
rows = []
for pl in used:
    for i, (v, en, de) in enumerate(PL[pl]):
        rows.append([f"{PFX} {pl}", v, en, de, 10 + i, "TRUE"])
clear(ws, 3)
put(ws, 3, rows, 3, 6)
ws.freeze_panes = "A3"

# Product to Product Constraints / Bundle Products / Non-config pricing
ws = wb["Prod-to-Prod Constraints"]; clear(ws, 2)
put(ws, 2, [[f"— No product-to-product constraints defined for {PNAME} —"]], 2, 4)
ws = wb["Bundle Products"]; clear(ws, 2)
put(ws, 2, [[f"— No bundle products defined for {PNAME} —"]], 2, 4)
ws = wb["Non-Config Product Pricing"]; clear(ws, 3)
put(ws, 3, [[f"— {PNAME} ({PCODE}) is configurable – pricing is in the 'Configurable Product Pricing' tab —"]], 3, 9)

# Product Internal Constraint
ws = wb["Product Internal Constraint"]
rows = [[PCODE, DRIVER.get(dv, C("Louvre") if dv in louvres else "—"), dv, dep, allowed, d]
        for dv, dep, allowed, d in CONSTRAINTS]
clear(ws, 2)
put(ws, 2, rows, 2, 6)

# Configurable Product Pricing (price request list – no prices in configurator)
ws = wb["Configurable Product Pricing"]
note_style = [copy.copy(ws.cell(3, c)._style) for c in range(1, 19)]
data_style_row = 3
clear(ws, 3)
for c in range(1, 19):
    ws.cell(3, c)._style = note_style[c - 1]
ws.merge_cells("A3:R3")
ws["A3"] = ("⚠ PRICES MUST BE REQUESTED FROM THE FACTORY — the Airlite configurator contains no prices "
            "(\"Prices available in ProGenTo or in the official price lists\" / \"Prices must be requested at the factory\"). "
            "Fill in the yellow cells: Global Cost (CC) and Product Selling Margin are mandatory.")
PRICED = ["Version", "Length", "Louvre", "Controls", "Thermal Release", "Control Bracket", "Base Design",
          "Flange Type", "Coating", "Coating Thickness", "Wind Baffle", "Delivery"]
code_of = {a[6]: C(a[7]) for a in CFG_ATTRS if a[6]}
rows = [[PNAME, PCODE, "Base price", "—", None, None, None, None, None, "Amount", None, None, None, None,
         "One Time", "TRUE", "EUR", "Base unit price; confirm with factory whether it is per unit or per m² (Av = Width × Length)"],
        [PNAME, PCODE, C("Width"), "per mm / per m²", None, None, None, None, None, "Glazing_Formula", None, None, None, None,
         "One Time", "TRUE", "EUR", "Size-dependent price (Width is a free number 500–2000 mm) – method to be confirmed by factory"]]
for pl in PRICED:
    for v, en, _ in PL[pl]:
        rows.append([PNAME, PCODE, code_of[pl], v, None, None, None, None, None, "Amount", None, None, None, None,
                     "One Time", "TRUE", "EUR", en])
font = Font(name="Arial", size=10)
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        cell = ws.cell(4 + i, j + 1, v)
        cell.font = font
        if j in (5, 13):  # Global Cost, Product Selling Margin
            cell.fill = YELLOW
ws.freeze_panes = "A4"

# Notes sheet (assumptions / open points)
ws = wb.create_sheet("Airlite Notes")
notes = [
    ("Airlite Tech. Configurator → RCA Product Price Template", ""),
    ("Source", "Airlite Tech. Configurator(1).xlsm, Version 1.54 (GPTS 1032, 2021-06), sheet Input + GPNDS description table (Input!BF:BI)"),
    ("Structure", "Same 10 sheets / columns as the KLAW Germany Product Price Template (ECO CO reference)"),
    ("Prefix / product code", f"{PFX} / {PCODE}"),
    ("Configurable attributes", f"{len(CFG_ATTRS)} (Input rows 13–37)"),
    ("Calculated output attributes", f"{len(OUT_ATTRS)} (Input rows 42–54) – read-only, category 'Calculated Outputs'"),
    ("Picklists / values", f"{len(used)} picklists, {sum(len(PL[p]) for p in used)} values"),
    ("Internal constraints", f"{len(CONSTRAINTS)} rules from Input columns AH (1E) / AJ (1S) with messages from column AM"),
    ("", ""),
    ("Open points / assumptions", ""),
    ("Prices", "No prices in the configurator. 'Configurable Product Pricing' is pre-filled as a price-request list (yellow cells to fill)."),
    ("Width", "Free number (500–2000 mm) in the configurator, so modelled as Number, not Picklist. Pricing method for size to be confirmed."),
    ("Length", "Dropdown allows 1200–3600 mm, but the validation message says max. 3000 mm. Confirm 3300/3600 with factory."),
    ("Controls M2B230", "Has a description in GPNDS but is NOT in the configurator dropdown – excluded."),
    ("Insulation", "Only value 'I' is selectable in the configurator."),
    ("Version 1V", "Referenced in some formulas but not selectable – ignored."),
    ("Snow / wind / CO2 checks", "Depend on large calculation sheets (SL Pneumatic Cylinders, SL Electric Drive, Windload). "
     "Need pre-computed CML tables or an Apex action, as with ECO."),
    ("External link", "Configurator links to C:\\Users\\deabomoh\\Desktop\\MONO as Example.xlsm (not needed for this template)."),
    ("To fill before go-live", "Product owner, availability date, prices, Product_Family, Pricebook."),
]
for i, (a, b) in enumerate(notes, 1):
    ws.cell(i, 1, a).font = Font(name="Arial", size=10, bold=(i == 1 or a == "Open points / assumptions"))
    ws.cell(i, 2, b).font = Font(name="Arial", size=10)
    ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(i, 1).alignment = Alignment(vertical="top")
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 110

wb.active = 0
wb.save(out_path)
print("attrs", len(CFG_ATTRS), "+", len(OUT_ATTRS), "| picklists", len(used),
      "| values", sum(len(PL[p]) for p in used), "| constraints", len(CONSTRAINTS), "| price rows", len(rows))
